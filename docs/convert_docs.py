import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

def add_markdown_to_doc(doc, filename):
    """Parses a markdown file and adds styled content to the docx object."""
    with open(filename, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Add a title page break for new files
    doc.add_page_break()
    doc.add_heading(f"Source: {filename}", 0)

    in_code_block = False
    code_buffer = []

    for line in lines:
        line = line.rstrip()
        
        # Handle Code Blocks (including Mermaid)
        if line.startswith("```"):
            if in_code_block:
                # End of code block - write buffer
                p = doc.add_paragraph()
                runner = p.add_run("\n".join(code_buffer))
                runner.font.name = 'Courier New'
                runner.font.size = Pt(9)
                runner.font.color.rgb = RGBColor(50, 50, 50) # Dark Grey
                p.style = 'Quote' # Use a distinct style for code
                in_code_block = False
                code_buffer = []
            else:
                # Start of code block
                in_code_block = True
            continue
        
        if in_code_block:
            code_buffer.append(line)
            continue

        # Handle Headers
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
        
        # Handle Lists
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            doc.add_paragraph(line.strip()[2:], style='List Bullet')
        elif line.strip().startswith("1. "):
            doc.add_paragraph(line.strip()[3:], style='List Number')
            
        # Handle Tables (Simple rendering)
        elif "|" in line:
            p = doc.add_paragraph(line)
            p.runs[0].font.name = 'Courier New'
            p.runs[0].font.size = Pt(9)
            
        # Standard Text
        else:
            if line.strip(): # Skip empty lines
                doc.add_paragraph(line)

def create_master_doc():
    doc = Document()
    
    # Add Title Page
    doc.add_heading('FraudGuard System Documentation', 0)
    doc.add_paragraph('Compiled Technical Reference')
    
    # List of your files in desired order
    files = [
        'README.md',
        'TECHNICAL_DOCUMENTATION.md',
        'ARCHITECTURE_DIAGRAMS.md',
        'ML_PIPELINE_ALGORITHMS.md'
    ]
    
    for file in files:
        if os.path.exists(file):
            print(f"Processing {file}...")
            add_markdown_to_doc(doc, file)
        else:
            print(f"Warning: {file} not found.")

    output_filename = 'FraudGuard_Full_Documentation.docx'
    doc.save(output_filename)
    print(f"Successfully created {output_filename}")

if __name__ == "__main__":
    create_master_doc()